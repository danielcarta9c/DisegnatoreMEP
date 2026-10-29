"""Computo per la quotazione a corpo dell'installatore (AA-01, tavole M.02 e M.03).
Stessi numeri delle tavole (dim.json). Solo valori, nessuna formula."""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

L = Path(__file__).parent
dim = json.loads((L / "dim.json").read_text(encoding="utf-8"))
REV = sys.argv[1] if len(sys.argv) > 1 else "rev00"
OUT = L / f"AA-01_M.02_Computo_canalizzato_per_installatore_{REV}.xlsx"

F = "Arial"
f_norm = Font(name=F, size=9)
f_bold = Font(name=F, size=9, bold=True)
f_tit = Font(name=F, size=12, bold=True, color="1C2E4A")
f_sub = Font(name=F, size=9, italic=True, color="555555")
f_avv = Font(name=F, size=9, bold=True, color="0000FF")
f_input = Font(name=F, size=9, color="0000FF")
f_head = Font(name=F, size=9, bold=True, color="FFFFFF")
f_grigio = Font(name=F, size=8, italic=True, color="7F7F7F")
fill_head = PatternFill("solid", fgColor="1C2E4A")
fill_grp = PatternFill("solid", fgColor="E6E9EF")
fill_in = PatternFill("solid", fgColor="FFFF00")
fill_na = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
EURO = '#,##0.00 "€"'

wb = Workbook()
ws = wb.active
ws.title = "Computo"
ws.sheet_view.showGridLines = False
testa = [
    ("AA-01 — Caso di prova aria-aria — Climatizzazione canalizzata Daikin + Airzone (dati del committente omessi)", f_tit),
    ("Computo allegato alle tavole M.02 (schema funzionale) e M.03 (pianta dei tracciati), rev. 00 del 29/10/2026.", f_sub),
    ("Offerta richiesta: PREZZO A CORPO, FISSO E INVARIABILE, per l'impianto finito e funzionante. Le quantità sono "
     "indicative e già maggiorate: l'installatore le verifica nel sopralluogo obbligatorio prima dell'offerta e ne "
     "risponde; quanto manca va aggiunto nella sezione H.", f_avv),
    ("Si compilano solo le celle gialle, in € oltre IVA: prezzo unitario e importo per le sezioni B-F e H, importo a "
     "corpo per la sezione G, solo prezzo unitario per le varianti V (fuori dal totale). Formato: 9,50. Righe grigie: "
     "fornite da Nove C o del Committente, non si quotano.", f_avv),
]
for i, (t, f) in enumerate(testa, start=1):
    ws.cell(i, 1, t).font = f
    ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=8)
    ws.cell(i, 1).alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[3].height = 26
ws.row_dimensions[4].height = 26

cols = [("Rif.", 7), ("Voce", 72), ("UM", 9), ("Q.tà", 8), ("Fornitura", 13), ("Note", 40),
        ("Prezzo unitario €", 14), ("Importo €", 14)]
R0 = 6
for j, (h, w) in enumerate(cols, start=1):
    c = ws.cell(R0, j, h)
    c.font, c.fill, c.border = f_head, fill_head, box
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(j)].width = w
ws.row_dimensions[R0].height = 26
r = R0 + 1


def gruppo(t):
    global r
    ws.cell(r, 1, t).font = f_bold
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    for j in range(1, 9):
        ws.cell(r, j).border = box
        ws.cell(r, j).fill = fill_grp
    r += 1


def riga(vals, quotare=True, solo_pu=False, solo_importo=False, motivo=""):
    global r
    for j, x in enumerate(vals, start=1):
        c = ws.cell(r, j, x)
        c.font, c.border, c.alignment = f_norm, box, wrap
    ws.cell(r, 4).alignment = Alignment(horizontal="right", vertical="top")
    pu, imp = ws.cell(r, 7), ws.cell(r, 8)
    pu.border = imp.border = box
    pu.number_format = imp.number_format = EURO
    if not quotare:
        pu.fill = imp.fill = fill_na
        imp.value, imp.font = motivo, f_grigio
    else:
        if not solo_importo:
            pu.fill, pu.font = fill_in, f_input
        else:
            pu.fill = fill_na
        if not solo_pu:
            imp.fill, imp.font = fill_in, f_input
        else:
            imp.fill = fill_na
            imp.value, imp.font = "variante, fuori totale", f_grigio
    r += 1


cat, n = None, {}
for v in dim["bom"]:
    if v["cat"] != cat:
        cat = v["cat"]
        gruppo(cat)
    lett = cat[0]
    n[lett] = n.get(lett, 0) + 1
    vals = [f"{lett}.{n[lett]}", v["voce"], v["um"], v["q"], v["fornitura"], v["note"]]
    if v["fornitura"] == "installatore":
        riga(vals)
    else:
        riga(vals, quotare=False, motivo="fornitura Nove C" if v["fornitura"] == "Nove C" else "escluso")

gruppo("G. Attività da quotare a corpo (manodopera, mezzi, prove, documenti)")
for i, a in enumerate(dim["attivita"], start=1):
    riga([f"G.{i}", a, "a corpo", 1, "installatore", ""], solo_importo=True)

gruppo("H. Voci aggiunte dall'installatore dopo il sopralluogo (necessarie per l'impianto finito)")
for i in range(1, 7):
    riga([f"H.{i}", "", "", "", "installatore", ""])

r += 1
ws.cell(r, 6, "TOTALE A CORPO, oltre IVA (somma degli importi B-H)").font = f_bold
ws.cell(r, 6).alignment = Alignment(horizontal="right", wrap_text=True)
t = ws.cell(r, 8)
t.fill, t.font, t.border, t.number_format = fill_in, Font(name=F, size=9, bold=True, color="0000FF"), box, EURO
r += 2

gruppo("V. Prezzi unitari per varianti (fuori dal prezzo a corpo, solo se Nove C le ordina per iscritto)")
for rif, voce_v, um in dim["varianti"]:
    riga([rif, voce_v, um, "", "installatore", ""], solo_pu=True)

r += 1
ws.cell(r, 1, "Validità dell'offerta, tempi di esecuzione (posa prima della chiusura del controsoffitto, messa in "
              "servizio al secondo accesso) e referente tecnico dell'installatore:").font = f_bold
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
r += 1
for j in range(1, 9):
    ws.cell(r, j).border = box
    ws.cell(r, j).fill = fill_in
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
ws.row_dimensions[r].height = 40

ws.freeze_panes = ws.cell(R0 + 1, 1)
ws.print_title_rows = f"{R0}:{R0}"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

# ------------------------------------------------------------------ tratte
wt = wb.create_sheet("Tratte")
wt.sheet_view.showGridLines = False
wt["A1"] = "Tabella delle tratte — diametri e lunghezze (tavola M.02; tracciati sulla tavola M.03)"
wt["A1"].font = f_tit
hdr = [("Tratto", 8), ("Rete", 20), ("Da", 9), ("A", 14), ("Ø / sezione / cavo", 18), ("Materiale", 30),
       ("Portata m³/h", 11), ("Velocità m/s", 11), ("L stimata m", 11), ("L da quotare m", 13)]
for j, (h, w) in enumerate(hdr, start=1):
    c = wt.cell(3, j, h)
    c.font, c.fill, c.border = f_head, fill_head, box
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wt.column_dimensions[get_column_letter(j)].width = w
T = {t["id"]: t for t in dim["tratte"]}
rr = 4
for tid in dim["ordine"]:
    t = T[tid]
    vals = [t["num"], t["rete"], t["da"].replace("QE-01", "quadro"), t["a"], t["dn"] or "", t["materiale"],
            t["portata_m3h"], t["v_ms"], t["l_stima_m"], t["l_quot_m"]]
    for j, x in enumerate(vals, start=1):
        c = wt.cell(rr, j, x)
        c.font, c.border = f_norm, box
        if j >= 7:
            c.number_format = "0" if j == 7 else "0.0"
    rr += 1
wt.cell(rr + 1, 1, "L da quotare = (L stimata + 1 m) × 1,15, per eccesso al mezzo metro; linee frigorifere 12 m "
                   "e condensa 10 m per l'incertezza dei percorsi.").font = f_sub
wt.freeze_panes = "A4"

# ------------------------------------------------------------------ note
wn = wb.create_sheet("Note")
wn.sheet_view.showGridLines = False
wn["A1"] = "Note tecniche (tavole M.02 e M.03)"
wn["A1"].font = f_tit
wn.column_dimensions["A"].width = 5
wn.column_dimensions["B"].width = 120
for i, t in enumerate(dim["note"], start=1):
    wn.cell(i + 2, 1, i).font = f_norm
    c = wn.cell(i + 2, 2, t)
    c.font, c.alignment = f_norm, wrap

wb.save(OUT)
print(OUT)
